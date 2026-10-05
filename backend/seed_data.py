"""Generate deterministic sample salary data for SQLite, CSV, or Excel."""

from __future__ import annotations

import argparse
import csv
import random
import re
import sqlite3
import unicodedata
from contextlib import closing
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import NamedTuple

from faker import Faker
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.database import connect_database, database_path, initialize_database
from backend.fx_reference import seed_usd_reference_rates


DEFAULT_AS_OF = date(2026, 10, 1)
DEFAULT_SEED = 42
EMPLOYEE_COUNT = 50


class CountryProfile(NamedTuple):
    code: str
    name: str
    currency_code: str
    currency_name: str
    symbol: str
    decimal_places: int
    locale: str
    locations: tuple[tuple[str, str, str], tuple[str, str, str]]
    salary_bands: dict[str, tuple[int, int]]


class DepartmentProfile(NamedTuple):
    code: str
    name: str
    title: str
    salary_percent: int


COUNTRIES = (
    CountryProfile(
        "IN", "India", "INR", "Indian Rupee", "₹", 2, "en_IN",
        (("Mumbai Office", "Mumbai", "Maharashtra"),
         ("Bengaluru Office", "Bengaluru", "Karnataka")),
        {"associate": (600_000, 1_200_000), "senior": (1_200_000, 2_400_000),
         "head": (3_500_000, 5_000_000)},
    ),
    CountryProfile(
        "US", "United States", "USD", "US Dollar", "$", 2, "en_US",
        (("New York Office", "New York", "New York"),
         ("Austin Office", "Austin", "Texas")),
        {"associate": (55_000, 85_000), "senior": (85_000, 130_000),
         "head": (170_000, 230_000)},
    ),
    CountryProfile(
        "GB", "United Kingdom", "GBP", "Pound Sterling", "£", 2, "en_GB",
        (("London Office", "London", "England"),
         ("Manchester Office", "Manchester", "England")),
        {"associate": (35_000, 55_000), "senior": (55_000, 85_000),
         "head": (110_000, 160_000)},
    ),
    CountryProfile(
        "DE", "Germany", "EUR", "Euro", "€", 2, "de_DE",
        (("Berlin Office", "Berlin", "Berlin"),
         ("Munich Office", "Munich", "Bavaria")),
        {"associate": (45_000, 65_000), "senior": (65_000, 100_000),
         "head": (130_000, 180_000)},
    ),
    CountryProfile(
        "SG", "Singapore", "SGD", "Singapore Dollar", "S$", 2, "en_GB",
        (("Central Office", "Singapore", "Central Region"),
         ("Jurong Office", "Singapore", "West Region")),
        {"associate": (55_000, 85_000), "senior": (85_000, 130_000),
         "head": (170_000, 230_000)},
    ),
)

DEPARTMENTS = (
    DepartmentProfile("HR", "Human Resources", "HR Specialist", 100),
    DepartmentProfile("ENG", "Engineering", "Software Engineer", 110),
    DepartmentProfile("FIN", "Finance", "Financial Analyst", 105),
    DepartmentProfile("OPS", "Operations", "Operations Specialist", 95),
    DepartmentProfile("SALES", "Sales", "Account Executive", 100),
)

ALLOWANCE_TYPES = (
    (1, "HOUSING", "Housing", "Monthly housing support", 7),
    (2, "TRANSPORT", "Transport", "Monthly transport support", 2),
    (3, "MEAL", "Meal", "Monthly meal support", 1),
    (4, "INTERNET", "Internet", "Monthly internet support", 1),
)

TABLE_COLUMNS = {
    "currency": ("currency_id", "currency_code", "currency_name", "symbol",
                 "decimal_places", "created_at", "updated_at"),
    "country": ("country_id", "country_code", "country_name", "default_currency_id",
                "created_at", "updated_at"),
    "location": ("location_id", "country_id", "location_name", "city", "state",
                 "created_at", "updated_at"),
    "department": ("department_id", "department_code", "department_name", "manager_id",
                   "created_at", "updated_at"),
    "employee": ("employee_id", "employee_code", "first_name", "last_name", "email",
                 "department_id", "location_id", "manager_id", "job_title",
                 "employment_type", "joining_date", "termination_date", "status",
                 "created_at", "updated_at"),
    "employee_compensation": ("id", "employee_id", "base_pay", "variable_pay",
                              "currency_id", "pay_frequency", "effective_from",
                              "effective_to", "created_at", "updated_at"),
    "allowance_type": ("id", "code", "name", "description", "created_at", "updated_at"),
    "employee_allowance": ("id", "compensation_id", "allowance_type_id", "amount",
                           "frequency", "created_at", "updated_at"),
}

Dataset = dict[str, list[dict[str, object]]]


def _timestamp(day: date) -> str:
    return f"{day.isoformat()} 12:00:00"


def _round_to_step(value: int, step: int) -> int:
    return ((value + step // 2) // step) * step


def _email_part(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", ".", ascii_value.lower()).strip(".") or "employee"


def generate_dataset(seed: int = DEFAULT_SEED, as_of: date = DEFAULT_AS_OF) -> Dataset:
    """Return 50 related rows using fixed IDs, local salaries, and seeded Faker names."""
    if as_of.year < 2010:
        raise ValueError("as_of must be in 2010 or later")

    rng = random.Random(seed)
    fakers = [Faker(country.locale) for country in COUNTRIES]
    for index, fake in enumerate(fakers):
        fake.seed_instance(seed + (index + 1) * 1009)

    dataset: Dataset = {table: [] for table in TABLE_COLUMNS}
    reference_stamp = _timestamp(as_of - timedelta(days=5 * 365))
    as_of_stamp = _timestamp(as_of)

    for country_id, country in enumerate(COUNTRIES, start=1):
        dataset["currency"].append({
            "currency_id": country_id,
            "currency_code": country.currency_code,
            "currency_name": country.currency_name,
            "symbol": country.symbol,
            "decimal_places": country.decimal_places,
            "created_at": reference_stamp,
            "updated_at": reference_stamp,
        })
        dataset["country"].append({
            "country_id": country_id,
            "country_code": country.code,
            "country_name": country.name,
            "default_currency_id": country_id,
            "created_at": reference_stamp,
            "updated_at": reference_stamp,
        })
        for location_index, (name, city, state) in enumerate(country.locations):
            dataset["location"].append({
                "location_id": (country_id - 1) * 2 + location_index + 1,
                "country_id": country_id,
                "location_name": name,
                "city": city,
                "state": state,
                "created_at": reference_stamp,
                "updated_at": reference_stamp,
            })

    # The first 25 employees cover every country/department pair once.
    manager_by_department = {
        department_index + 1: department_index * 5 + department_index + 1
        for department_index in range(len(DEPARTMENTS))
    }
    for department_id, department in enumerate(DEPARTMENTS, start=1):
        dataset["department"].append({
            "department_id": department_id,
            "department_code": department.code,
            "department_name": department.name,
            "manager_id": manager_by_department[department_id],
            "created_at": reference_stamp,
            "updated_at": as_of_stamp,
        })

    for allowance_id, code, name, description, _percent in ALLOWANCE_TYPES:
        dataset["allowance_type"].append({
            "id": allowance_id,
            "code": code,
            "name": name,
            "description": description,
            "created_at": reference_stamp,
            "updated_at": reference_stamp,
        })

    compensation_id = 1
    employee_allowance_id = 1
    for index in range(EMPLOYEE_COUNT):
        employee_id = index + 1
        country_index = index % len(COUNTRIES)
        department_index = (index // len(COUNTRIES)) % len(DEPARTMENTS)
        country = COUNTRIES[country_index]
        department = DEPARTMENTS[department_index]
        department_id = department_index + 1
        head_id = manager_by_department[department_id]
        level = "head" if employee_id == head_id else rng.choices(
            ("associate", "senior"), weights=(6, 4), k=1
        )[0]

        fake = fakers[country_index]
        first_name, last_name = fake.first_name(), fake.last_name()
        employee_code = f"EMP{employee_id:04d}"
        email = f"{_email_part(first_name)}.{_email_part(last_name)}.{employee_code.lower()}@acme.org"
        joining_date = as_of - timedelta(days=rng.randint(3 * 365, 8 * 365))
        current_start = as_of - timedelta(days=rng.randint(60, 420))
        prior_start = max(
            joining_date,
            current_start - timedelta(days=rng.randint(300, 700)),
        )
        prior_end = current_start - timedelta(days=1)
        title = f"Head of {department.name}" if level == "head" else (
            f"Senior {department.title}" if level == "senior" else department.title
        )
        dataset["employee"].append({
            "employee_id": employee_id,
            "employee_code": employee_code,
            "first_name": first_name,
            "last_name": last_name,
            "email": email,
            "department_id": department_id,
            "location_id": country_index * 2 + (index // len(COUNTRIES)) % 2 + 1,
            "manager_id": None if employee_id == head_id else head_id,
            "job_title": title,
            "employment_type": "FULL_TIME",
            "joining_date": joining_date.isoformat(),
            "termination_date": None,
            "status": "ACTIVE",
            "created_at": _timestamp(joining_date),
            "updated_at": as_of_stamp,
        })

        low, high = country.salary_bands[level]
        base_major = rng.randrange(low // 1000, high // 1000 + 1) * 1000
        base_major = _round_to_step(base_major * department.salary_percent // 100, 1000)
        scale = 10 ** country.decimal_places
        current_base = base_major * scale
        raise_percent = rng.randint(7, 16)
        prior_base = _round_to_step(current_base * 100 // (100 + raise_percent), 1000 * scale)

        variable_percent: int | None = None
        if department.code == "SALES":
            variable_percent = rng.randint(8, 18)
        elif level == "head":
            variable_percent = rng.randint(5, 12)
        elif department.code == "FIN" and level == "senior":
            variable_percent = rng.randint(3, 7)

        allowance_type_ids = sorted(rng.sample(range(1, len(ALLOWANCE_TYPES) + 1), rng.randint(1, 3)))
        for base_pay, effective_from, effective_to, created_at, updated_at in (
            (prior_base, prior_start, prior_end, _timestamp(prior_start), _timestamp(current_start)),
            (current_base, current_start, None, _timestamp(current_start), as_of_stamp),
        ):
            variable_pay = None if variable_percent is None else _round_to_step(
                base_pay * variable_percent // 100, 100 * scale
            )
            dataset["employee_compensation"].append({
                "id": compensation_id,
                "employee_id": employee_id,
                "base_pay": base_pay,
                "variable_pay": variable_pay,
                "currency_id": country_index + 1,
                "pay_frequency": "ANNUAL",
                "effective_from": effective_from.isoformat(),
                "effective_to": effective_to.isoformat() if effective_to else None,
                "created_at": created_at,
                "updated_at": updated_at,
            })
            for allowance_type_id in allowance_type_ids:
                percentage = ALLOWANCE_TYPES[allowance_type_id - 1][4]
                percentage = max(1, percentage + rng.randint(-1, 1))
                monthly_amount = max(
                    10 * scale,
                    _round_to_step(base_pay * percentage // 1200, 10 * scale),
                )
                dataset["employee_allowance"].append({
                    "id": employee_allowance_id,
                    "compensation_id": compensation_id,
                    "allowance_type_id": allowance_type_id,
                    "amount": monthly_amount,
                    "frequency": "MONTHLY",
                    "created_at": created_at,
                    "updated_at": updated_at,
                })
                employee_allowance_id += 1
            compensation_id += 1

    return dataset


def _insert_rows(connection: sqlite3.Connection, table: str, rows: list[dict[str, object]]) -> None:
    columns = TABLE_COLUMNS[table]
    column_names = ", ".join(columns)
    placeholders = ", ".join(f":{column}" for column in columns)
    connection.executemany(
        f"INSERT INTO {table} ({column_names}) VALUES ({placeholders})",
        rows,
    )


def seed_sqlite(dataset: Dataset, path: Path) -> Path:
    """Insert the generated dataset atomically into an empty business schema."""
    initialize_database(path)
    with closing(connect_database(path)) as connection:
        with connection:
            connection.execute("BEGIN IMMEDIATE")
            populated = [
                table for table in TABLE_COLUMNS
                if connection.execute(f"SELECT 1 FROM {table} LIMIT 1").fetchone()
            ]
            if populated:
                raise ValueError(
                    "Business tables already contain data: " + ", ".join(populated)
                )

            for table in ("currency", "country", "location"):
                _insert_rows(connection, table, dataset[table])
            seed_usd_reference_rates(connection)

            departments = [
                row | {"manager_id": None, "updated_at": row["created_at"]}
                for row in dataset["department"]
            ]
            _insert_rows(connection, "department", departments)

            employees = [
                row | {"manager_id": None, "updated_at": row["created_at"]}
                if row["manager_id"] is not None else row
                for row in dataset["employee"]
            ]
            _insert_rows(connection, "employee", employees)
            connection.executemany(
                "UPDATE employee SET manager_id = :manager_id, updated_at = :updated_at "
                "WHERE employee_id = :employee_id",
                [row for row in dataset["employee"] if row["manager_id"] is not None],
            )
            connection.executemany(
                "UPDATE department SET manager_id = :manager_id, updated_at = :updated_at "
                "WHERE department_id = :department_id",
                dataset["department"],
            )

            for table in ("employee_compensation", "allowance_type", "employee_allowance"):
                _insert_rows(connection, table, dataset[table])

            foreign_key_issues = connection.execute("PRAGMA foreign_key_check").fetchall()
            if foreign_key_issues:
                raise ValueError(f"Foreign-key check failed: {foreign_key_issues}")

    return path


def write_csv(dataset: Dataset, directory: Path) -> Path:
    """Write one UTF-8 CSV file per generated table."""
    directory.mkdir(parents=True, exist_ok=False)
    for table, columns in TABLE_COLUMNS.items():
        with (directory / f"{table}.csv").open("w", newline="", encoding="utf-8") as output:
            writer = csv.DictWriter(output, fieldnames=columns)
            writer.writeheader()
            writer.writerows(dataset[table])
    return directory


def write_xlsx(dataset: Dataset, path: Path) -> Path:
    """Write the same related rows to a filterable Excel workbook."""
    if path.exists():
        raise FileExistsError(f"Output already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "ACME synthetic salary seed data"
    header_fill = PatternFill("solid", fgColor="1F2937")

    for table, columns in TABLE_COLUMNS.items():
        sheet = workbook.create_sheet(table)
        sheet.append(columns)
        for row in dataset[table]:
            values = []
            for column in columns:
                value = row[column]
                if isinstance(value, str) and column in {
                    "joining_date", "termination_date", "effective_from", "effective_to"
                }:
                    value = date.fromisoformat(value)
                elif isinstance(value, str) and column in {"created_at", "updated_at"}:
                    value = datetime.fromisoformat(value)
                values.append(value)
            sheet.append(values)

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 22
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for column_index, column in enumerate(columns, start=1):
            width = min(34, max(len(column), *(len(str(row[column] or "")) for row in dataset[table])) + 2)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            if column in {"joining_date", "termination_date", "effective_from", "effective_to"}:
                number_format = "yyyy-mm-dd"
            elif column in {"created_at", "updated_at"}:
                number_format = "yyyy-mm-dd hh:mm:ss"
            else:
                continue
            for cells in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                for cell in cells:
                    cell.number_format = number_format

    workbook.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--format", required=True, choices=("sqlite", "csv", "xlsx"))
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--as-of", type=date.fromisoformat, default=DEFAULT_AS_OF)
    parser.add_argument("--output", type=Path, help="CSV directory or XLSX file")
    parser.add_argument("--db-path", type=Path, help="SQLite file; defaults to DB_PATH")
    arguments = parser.parse_args()

    if arguments.format == "sqlite" and arguments.output is not None:
        parser.error("--output is only for csv or xlsx; use --db-path for sqlite")
    if arguments.format != "sqlite" and arguments.db_path is not None:
        parser.error("--db-path is only for sqlite")

    try:
        dataset = generate_dataset(arguments.seed, arguments.as_of)
        if arguments.format == "sqlite":
            path = (arguments.db_path.expanduser().resolve() if arguments.db_path
                    else database_path())
            destination = seed_sqlite(dataset, path)
        elif arguments.format == "csv":
            destination = write_csv(
                dataset,
                (arguments.output or Path("exports/seed-data")).expanduser().resolve(),
            )
        else:
            destination = (arguments.output or Path("exports/seed-data.xlsx")).expanduser().resolve()
            if destination.suffix.lower() != ".xlsx":
                parser.error("Excel output must have an .xlsx extension")
            write_xlsx(dataset, destination)
    except (ValueError, OSError, sqlite3.Error) as error:
        parser.exit(1, f"Seed generation failed: {error}\n")

    print(
        f"Generated {len(dataset['employee'])} employees, "
        f"{len(dataset['employee_compensation'])} compensation records, "
        f"and {len(dataset['employee_allowance'])} allowances: {destination}"
    )


if __name__ == "__main__":
    main()
