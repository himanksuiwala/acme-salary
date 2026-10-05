"""Exercise deterministic generation and each output format."""

import csv
import tempfile
import unittest
from collections import Counter, defaultdict
from contextlib import closing
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from backend.database import connect_database
from backend.seed_data import generate_dataset, seed_sqlite, write_csv, write_xlsx


class SeedDataTests(unittest.TestCase):
    def test_generated_rows_are_deterministic_and_related(self) -> None:
        dataset = generate_dataset()
        self.assertEqual(dataset, generate_dataset())
        self.assertNotEqual(dataset["employee"], generate_dataset(seed=43)["employee"])
        self.assertEqual(len(dataset["employee"]), 10_000)
        self.assertEqual(len(dataset["employee_compensation"]), 20_000)

        location_country = {row["location_id"]: row["country_id"] for row in dataset["location"]}
        country_currency = {row["country_id"]: row["default_currency_id"] for row in dataset["country"]}
        country_counts = Counter(location_country[row["location_id"]] for row in dataset["employee"])
        self.assertEqual(sorted(country_counts.values()), [2_000] * 5)

        compensations = defaultdict(list)
        for row in dataset["employee_compensation"]:
            compensations[row["employee_id"]].append(row)
        allowances = defaultdict(list)
        for row in dataset["employee_allowance"]:
            allowances[row["compensation_id"]].append(row)

        for employee in dataset["employee"]:
            prior, current = compensations[employee["employee_id"]]
            currency_id = country_currency[location_country[employee["location_id"]]]
            self.assertEqual((prior["currency_id"], current["currency_id"]), (currency_id, currency_id))
            self.assertEqual(
                date.fromisoformat(prior["effective_to"]),
                date.fromisoformat(current["effective_from"]) - timedelta(days=1),
            )
            self.assertIsNone(current["effective_to"])
            self.assertLess(prior["base_pay"], current["base_pay"])
            for package in (prior, current):
                package_allowances = allowances[package["id"]]
                self.assertIn(len(package_allowances), (1, 2, 3))
                self.assertEqual(
                    len({row["allowance_type_id"] for row in package_allowances}),
                    len(package_allowances),
                )

    def test_employee_count_is_configurable(self) -> None:
        dataset = generate_dataset(employee_count=50)
        self.assertEqual(len(dataset["employee"]), 50)
        self.assertEqual(len(dataset["employee_compensation"]), 100)
        self.assertEqual(dataset["employee"][-1]["employee_code"], "EMP0050")

        with self.assertRaisesRegex(ValueError, "at least 25"):
            generate_dataset(employee_count=24)

    def test_sqlite_seeding_is_valid_and_refuses_a_second_seed(self) -> None:
        dataset = generate_dataset()
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "seed.db"
            seed_sqlite(dataset, path)
            with closing(connect_database(path)) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM employee").fetchone()[0], 10_000)
                self.assertEqual(
                    connection.execute("SELECT COUNT(*) FROM employee_compensation").fetchone()[0], 20_000
                )
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM app_user").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM audit_log").fetchone()[0], 0)
                self.assertEqual(connection.execute("PRAGMA foreign_key_check").fetchall(), [])
                self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")

            with self.assertRaisesRegex(ValueError, "already contain data"):
                seed_sqlite(dataset, path)

    def test_csv_and_excel_exports_keep_keys_and_typed_dates(self) -> None:
        dataset = generate_dataset(employee_count=50)
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            csv_directory = write_csv(dataset, root / "csv")
            self.assertEqual(len(list(csv_directory.glob("*.csv"))), 8)
            with (csv_directory / "employee.csv").open(newline="", encoding="utf-8") as file:
                employees = list(csv.DictReader(file))
            self.assertEqual(len(employees), 50)
            self.assertEqual(employees[0]["employee_code"], dataset["employee"][0]["employee_code"])

            workbook_path = write_xlsx(dataset, root / "seed.xlsx")
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, list(dataset))
                employee_sheet = workbook["employee"]
                self.assertEqual(employee_sheet.max_row - 1, 50)
                self.assertEqual(employee_sheet["B2"].value, dataset["employee"][0]["employee_code"])
                self.assertIsInstance(employee_sheet["K2"].value, datetime)
                compensation_sheet = workbook["employee_compensation"]
                self.assertEqual(compensation_sheet.max_row - 1, 100)
                self.assertIsInstance(compensation_sheet["C2"].value, int)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
